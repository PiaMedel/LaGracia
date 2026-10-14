"""
Carga masiva del catálogo desde UN archivo Excel con varias hojas:

  1. Materias primas
  2. Pre-preparaciones            (opcional)
  3. Productos
  4. Recetas pre-preparaciones    (opcional)
  5. Recetas productos

El orden de importación es fijo (las recetas dependen de que existan los
ingredientes y los productos). Es idempotente: si un nombre ya existe se
ACTUALIZA en vez de duplicarse, y las recetas se reemplazan completas.
Las filas cuyo primer valor empieza con "(ejemplo)" se ignoran.
"""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from extensions import db
from models.preparations import Preparation, PreparationRecipe, PreparationRecipeItem
from models.products import Product, ESTADOS_PRODUCTO
from models.raw_materials import RawMaterial
from models.recipes import Recipe, RecipeItem
from services.inventory_engine import registrar_movimiento

# ---------------------------------------------------------------- estructura
SHEETS = {
    "Materias primas": [
        "nombre", "unidad_medida", "costo_unitario", "stock_seguridad",
        "consumo_diario_estimado", "dias_reposicion", "compra_minima",
        "proveedor", "vida_util_dias", "stock_inicial",
    ],
    "Pre-preparaciones": [
        "nombre", "unidad_medida", "stock_seguridad",
        "vida_util_congelada_dias", "vida_util_descongelada_dias", "stock_inicial",
    ],
    "Productos": [
        "nombre", "categoria", "unidad_venta", "precio_venta", "estado",
        "stock_seguridad", "vida_util_dias", "stock_inicial",
    ],
    "Recetas pre-preparaciones": [
        "preparacion", "rendimiento", "materia_prima", "cantidad",
    ],
    "Recetas productos": [
        "producto", "rendimiento", "tipo_ingrediente", "ingrediente", "cantidad",
    ],
}

EXAMPLES = {
    "Materias primas": ["(ejemplo) Harina almendras", "kg", 6000, 2, 2, 3, 10, "Proveedor A", 180, 5],
    "Pre-preparaciones": ["(ejemplo) Masa base brownie", "porción", 6, 90, 2, 0],
    "Productos": ["(ejemplo) Brownie", "Bizcochería", "unidad", 4000, "POTENCIAL", 5, 5, 0],
    "Recetas pre-preparaciones": ["(ejemplo) Masa base brownie", 12, "Cacao", 0.1],
    "Recetas productos": ["(ejemplo) Brownie", 12, "Materia prima", "Harina almendras", 0.3],
}

LEGEND = [
    "CÓMO LLENAR ESTE ARCHIVO",
    "",
    "1. Llena las hojas en este orden: Materias primas → Pre-preparaciones (si usas) → Productos → Recetas.",
    "2. Cada hoja trae una fila de ejemplo que empieza con '(ejemplo)'. Puedes borrarla o dejarla: se ignora al importar.",
    "3. Los nombres deben escribirse IGUAL en todas las hojas (mayúsculas/minúsculas no importan). Ej: si en 'Materias primas' dice 'Harina almendras', en las recetas debe decir 'Harina almendras'.",
    "4. UNIDADES: la cantidad de una receta va en la MISMA unidad de la materia prima. Si la harina está en kg, 300 g se escribe 0.3 (o 0,3).",
    "5. 'rendimiento' = cuántas unidades produce UNA tanda de la receta (la cantidad de ingredientes es por esa tanda completa).",
    "6. 'tipo_ingrediente' en Recetas productos: escribe 'Materia prima' o 'Pre-preparación'.",
    "7. 'estado' de producto: GANADOR, POTENCIAL, REVISAR o PAUSADO (si lo dejas vacío queda POTENCIAL).",
    "8. 'stock_inicial': déjalo vacío o en 0. El stock real de hoy se fija al final, con la pantalla 'Conteo' de la app (así no se cuenta doble con el historial).",
    "9. Puedes volver a subir el archivo corregido: lo que ya existe se actualiza, no se duplica. Las recetas se reemplazan completas.",
    "10. Escribe los números como números (sin signo $ ni puntos de miles).",
]

LEGEND_SIMPLE = [
    "CÓMO LLENAR ESTE ARCHIVO",
    "",
    "1. Llena las hojas en este orden: Materias primas → Productos → Recetas productos.",
    "2. Cada hoja trae una fila de ejemplo que empieza con '(ejemplo)'. Puedes borrarla o dejarla: se ignora al importar.",
    "3. Los nombres deben escribirse IGUAL en todas las hojas (mayúsculas/minúsculas no importan). Ej: si en 'Materias primas' dice 'Harina almendras', en las recetas debe decir 'Harina almendras'.",
    "4. UNIDADES: la cantidad de una receta va en la MISMA unidad de la materia prima. Si la harina está en kg, 300 g se escribe 0.3 (o 0,3).",
    "5. 'rendimiento' = cuántas unidades produce UNA tanda de la receta (la cantidad de ingredientes es por esa tanda completa).",
    "6. Una receta ocupa varias filas: una por ingrediente, repitiendo el nombre del producto y el rendimiento.",
    "7. 'estado' de producto: GANADOR, POTENCIAL, REVISAR o PAUSADO (si lo dejas vacío queda POTENCIAL).",
    "8. 'stock_inicial': déjalo vacío o en 0. El stock real de hoy se fija al final, con la pantalla 'Conteo' de la app (así no se cuenta doble con el historial).",
    "9. Puedes volver a subir el archivo corregido: lo que ya existe se actualiza, no se duplica. Las recetas se reemplazan completas.",
    "10. Escribe los números como números (sin signo $ ni puntos de miles).",
]

SIMPLE_SHEETS = {
    "Materias primas": SHEETS["Materias primas"],
    "Productos": SHEETS["Productos"],
    "Recetas productos": ["producto", "rendimiento", "ingrediente", "cantidad"],
}
SIMPLE_EXAMPLES = {
    "Materias primas": EXAMPLES["Materias primas"],
    "Productos": EXAMPLES["Productos"],
    "Recetas productos": ["(ejemplo) Brownie", 12, "Harina almendras", 0.3],
}

_FONT = "Arial"


# ------------------------------------------------------------------ plantilla
def build_template_bytes(usar_preps=False):
    sheets = SHEETS if usar_preps else SIMPLE_SHEETS
    examples = EXAMPLES if usar_preps else SIMPLE_EXAMPLES
    legend = LEGEND if usar_preps else LEGEND_SIMPLE
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "LEEME"
    for i, line in enumerate(legend, start=1):
        c = ws0.cell(row=i, column=1, value=line)
        c.font = Font(name=_FONT, size=12 if i == 1 else 10, bold=(i == 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 130

    head_fill = PatternFill("solid", fgColor="A9583C")
    example_fill = PatternFill("solid", fgColor="FFF2CC")
    for sheet_name, cols in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for j, col in enumerate(cols, start=1):
            c = ws.cell(row=1, column=j, value=col)
            c.font = Font(name=_FONT, bold=True, color="FFFFFF")
            c.fill = head_fill
            ws.column_dimensions[c.column_letter].width = max(16, len(col) + 4)
        for j, val in enumerate(examples[sheet_name], start=1):
            c = ws.cell(row=2, column=j, value=val)
            c.font = Font(name=_FONT, italic=True)
            c.fill = example_fill
        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------- helpers
def _txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _num(v):
    """Número o None. Acepta coma decimal ('0,3') por si lo escriben como texto."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).strip().replace(" ", "").replace("$", "").replace(",", "."))


def _int(v):
    n = _num(v)
    return int(round(n)) if n is not None else None


def _read_sheet(wb, name):
    """Devuelve [(n_fila_excel, {col: valor}), ...] ignorando vacías y ejemplos."""
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    keys = [(_txt(h) or "").lower().replace(" ", "_") for h in header]
    out = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        first = _txt(row[0]) or ""
        if first.lower().startswith("(ejemplo)"):
            continue
        out.append((idx, dict(zip(keys, row))))
    return out


def _new_result():
    return {"creados": 0, "actualizados": 0, "errores": []}


def _set_if(obj, attr, value):
    if value is not None:
        setattr(obj, attr, value)


# ----------------------------------------------------------------- secciones
def _import_materials(rows, res):
    existing = {m.nombre.strip().lower(): m for m in RawMaterial.query.all()}
    for idx, r in rows:
        try:
            nombre = _txt(r.get("nombre"))
            if not nombre:
                raise ValueError("falta el nombre")
            m = existing.get(nombre.lower())
            nuevo = m is None
            if nuevo:
                m = RawMaterial(nombre=nombre)
                db.session.add(m)
            _set_if(m, "unidad_medida", _txt(r.get("unidad_medida")))
            _set_if(m, "costo_unitario", _num(r.get("costo_unitario")))
            _set_if(m, "stock_seguridad", _num(r.get("stock_seguridad")))
            _set_if(m, "consumo_diario_estimado", _num(r.get("consumo_diario_estimado")))
            _set_if(m, "dias_reposicion", _int(r.get("dias_reposicion")))
            _set_if(m, "compra_minima", _num(r.get("compra_minima")))
            _set_if(m, "proveedor", _txt(r.get("proveedor")))
            _set_if(m, "vida_util_dias", _int(r.get("vida_util_dias")))
            db.session.flush()
            stock = _num(r.get("stock_inicial"))
            if nuevo and stock:
                registrar_movimiento("material", m.id, stock, "stock_inicial", commit=False)
            existing[nombre.lower()] = m
            res["creados" if nuevo else "actualizados"] += 1
        except Exception as e:  # noqa: BLE001
            res["errores"].append(f"Materias primas, fila {idx}: {e}")


def _import_preparations(rows, res):
    existing = {p.nombre.strip().lower(): p for p in Preparation.query.all()}
    for idx, r in rows:
        try:
            nombre = _txt(r.get("nombre"))
            if not nombre:
                raise ValueError("falta el nombre")
            p = existing.get(nombre.lower())
            nuevo = p is None
            if nuevo:
                p = Preparation(nombre=nombre)
                db.session.add(p)
            _set_if(p, "unidad_medida", _txt(r.get("unidad_medida")))
            _set_if(p, "stock_seguridad", _num(r.get("stock_seguridad")))
            _set_if(p, "vida_util_congelada_dias", _int(r.get("vida_util_congelada_dias")))
            _set_if(p, "vida_util_descongelada_dias", _int(r.get("vida_util_descongelada_dias")))
            db.session.flush()
            stock = _num(r.get("stock_inicial"))
            if nuevo and stock:
                registrar_movimiento("preparation", p.id, stock, "stock_inicial", commit=False)
            existing[nombre.lower()] = p
            res["creados" if nuevo else "actualizados"] += 1
        except Exception as e:  # noqa: BLE001
            res["errores"].append(f"Pre-preparaciones, fila {idx}: {e}")


def _import_products(rows, res):
    existing = {p.nombre.strip().lower(): p for p in Product.query.all()}
    for idx, r in rows:
        try:
            nombre = _txt(r.get("nombre"))
            if not nombre:
                raise ValueError("falta el nombre")
            estado = (_txt(r.get("estado")) or "").upper() or None
            if estado and estado not in ESTADOS_PRODUCTO:
                raise ValueError(f"estado '{estado}' no válido (usa {', '.join(ESTADOS_PRODUCTO)})")
            p = existing.get(nombre.lower())
            nuevo = p is None
            if nuevo:
                p = Product(nombre=nombre, estado="POTENCIAL")
                db.session.add(p)
            _set_if(p, "categoria", _txt(r.get("categoria")))
            _set_if(p, "unidad_venta", _txt(r.get("unidad_venta")))
            _set_if(p, "precio_venta", _num(r.get("precio_venta")))
            _set_if(p, "estado", estado)
            _set_if(p, "stock_seguridad", _num(r.get("stock_seguridad")))
            _set_if(p, "vida_util_dias", _int(r.get("vida_util_dias")))
            db.session.flush()
            stock = _num(r.get("stock_inicial"))
            if nuevo and stock:
                registrar_movimiento("product", p.id, stock, "stock_inicial", commit=False)
            existing[nombre.lower()] = p
            res["creados" if nuevo else "actualizados"] += 1
        except Exception as e:  # noqa: BLE001
            res["errores"].append(f"Productos, fila {idx}: {e}")


def _group(rows, key):
    groups = {}
    for idx, r in rows:
        name = _txt(r.get(key))
        if name:
            groups.setdefault(name, []).append((idx, r))
        else:
            groups.setdefault(None, []).append((idx, r))
    return groups


def _import_prep_recipes(rows, res):
    preps = {p.nombre.strip().lower(): p for p in Preparation.query.all()}
    mats = {m.nombre.strip().lower(): m for m in RawMaterial.query.all()}
    for name, items in _group(rows, "preparacion").items():
        if name is None:
            for idx, _ in items:
                res["errores"].append(f"Recetas pre-preparaciones, fila {idx}: falta el nombre de la preparación")
            continue
        try:
            prep = preps.get(name.lower())
            if prep is None:
                raise ValueError(f"la preparación '{name}' no existe (créala en la hoja Pre-preparaciones)")
            rendimiento = next((_num(r.get("rendimiento")) for _, r in items if _num(r.get("rendimiento"))), None)
            if not rendimiento:
                raise ValueError("falta el rendimiento")
            new_items = []
            for idx, r in items:
                mname = _txt(r.get("materia_prima"))
                m = mats.get((mname or "").lower())
                if m is None:
                    raise ValueError(f"fila {idx}: materia prima '{mname}' no existe")
                cant = _num(r.get("cantidad"))
                if cant is None or cant <= 0:
                    raise ValueError(f"fila {idx}: cantidad inválida para '{mname}'")
                new_items.append(PreparationRecipeItem(raw_material_id=m.id, cantidad=cant))
            recipe = prep.recipe
            nuevo = recipe is None
            if nuevo:
                recipe = PreparationRecipe(preparation_id=prep.id, rendimiento=rendimiento)
                db.session.add(recipe)
            recipe.rendimiento = rendimiento
            recipe.items = new_items
            db.session.flush()
            res["creados" if nuevo else "actualizados"] += 1
        except Exception as e:  # noqa: BLE001
            res["errores"].append(f"Receta de pre-preparación '{name}' NO importada: {e}")


def _import_product_recipes(rows, res):
    prods = {p.nombre.strip().lower(): p for p in Product.query.all()}
    preps = {p.nombre.strip().lower(): p for p in Preparation.query.all()}
    mats = {m.nombre.strip().lower(): m for m in RawMaterial.query.all()}
    for name, items in _group(rows, "producto").items():
        if name is None:
            for idx, _ in items:
                res["errores"].append(f"Recetas productos, fila {idx}: falta el nombre del producto")
            continue
        try:
            prod = prods.get(name.lower())
            if prod is None:
                raise ValueError(f"el producto '{name}' no existe (créalo en la hoja Productos)")
            rendimiento = next((_num(r.get("rendimiento")) for _, r in items if _num(r.get("rendimiento"))), None)
            if not rendimiento:
                raise ValueError("falta el rendimiento")
            new_items = []
            for idx, r in items:
                tipo = (_txt(r.get("tipo_ingrediente")) or "materia prima").lower()
                ing = _txt(r.get("ingrediente"))
                cant = _num(r.get("cantidad"))
                if cant is None or cant <= 0:
                    raise ValueError(f"fila {idx}: cantidad inválida para '{ing}'")
                if tipo.startswith("pre"):
                    obj = preps.get((ing or "").lower())
                    if obj is None:
                        raise ValueError(f"fila {idx}: pre-preparación '{ing}' no existe")
                    new_items.append(RecipeItem(ingredient_type="preparation", preparation_id=obj.id, cantidad=cant))
                else:
                    obj = mats.get((ing or "").lower())
                    if obj is None:
                        raise ValueError(f"fila {idx}: materia prima '{ing}' no existe")
                    new_items.append(RecipeItem(ingredient_type="material", raw_material_id=obj.id, cantidad=cant))
            recipe = prod.recipe
            nuevo = recipe is None
            if nuevo:
                recipe = Recipe(product_id=prod.id, rendimiento=rendimiento)
                db.session.add(recipe)
            recipe.rendimiento = rendimiento
            recipe.items = new_items
            db.session.flush()
            res["creados" if nuevo else "actualizados"] += 1
        except Exception as e:  # noqa: BLE001
            res["errores"].append(f"Receta de '{name}' NO importada: {e}")


# ------------------------------------------------------------------ principal
def import_catalog(file_stream):
    """Importa el Excel completo. Devuelve {hoja: {creados, actualizados, errores}}."""
    wb = load_workbook(file_stream, data_only=True)
    results = {k: _new_result() for k in SHEETS}

    _import_materials(_read_sheet(wb, "Materias primas"), results["Materias primas"])
    db.session.commit()
    _import_preparations(_read_sheet(wb, "Pre-preparaciones"), results["Pre-preparaciones"])
    db.session.commit()
    _import_products(_read_sheet(wb, "Productos"), results["Productos"])
    db.session.commit()
    _import_prep_recipes(_read_sheet(wb, "Recetas pre-preparaciones"), results["Recetas pre-preparaciones"])
    db.session.commit()
    _import_product_recipes(_read_sheet(wb, "Recetas productos"), results["Recetas productos"])
    db.session.commit()
    return results
